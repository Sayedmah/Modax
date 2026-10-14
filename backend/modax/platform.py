"""User-scoped provider configuration, document retrieval, and measured evaluations."""

import asyncio
import base64
import io
import json
import os
import subprocess
import sys
import time
import zipfile
from pathlib import Path
from typing import Literal
from urllib.parse import urlparse

from cryptography.fernet import Fernet
from fastapi import Depends, HTTPException, UploadFile
from pydantic import BaseModel, Field

from .providers import HTTPProvider, Model, ProviderError
from .store import uid
from .workspace import SECRET

ALLOWED_HOSTS = {
    "api.openai.com",
    "api.anthropic.com",
    "openrouter.ai",
    "generativelanguage.googleapis.com",
}


class ProfileInput(BaseModel):
    family: Literal["gemini", "openrouter", "compatible", "anthropic"]
    model: str = Field(min_length=1, max_length=100)
    api_key: str = Field(min_length=6, max_length=512)
    base_url: str = Field(default="https://api.openai.com/v1", max_length=200)


class Preferences(BaseModel):
    use_memory: bool = False
    requests_per_day: int = Field(default=100, ge=1, le=500)


class ArenaInput(BaseModel):
    models: list[str] = Field(min_length=1, max_length=3)
    confirm_external_calls: bool = False
    template: str = Field(default="{task}", max_length=1000)


DATASET = [
    {"id": "arithmetic", "task": "Return only the numeric result of 7 * 8.", "expected": "56"},
    {
        "id": "instruction",
        "task": "Return exactly the uppercase word MODAX with no explanation.",
        "expected": "MODAX",
    },
    {
        "id": "json",
        "task": "Return only valid JSON with key status and value ok.",
        "expected": '{"status":"ok"}',
    },
]


def extract_document(filename, data):
    kind = Path(filename).suffix.lower()
    if len(data) > 200000:
        raise ValueError("File size limit is 200 KB")
    if kind in {".txt", ".md", ".csv", ".json", ".xml", ".log", ".py"}:
        text = data.decode("utf-8-sig")
        if kind == ".json":
            json.loads(text)
    elif kind in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 8000000:
                raise ValueError("Expanded document exceeds limit")
        if kind == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(
                "\t".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            )
        else:
            import openpyxl

            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            try:
                for sheet in book:
                    parts.append("Sheet: " + sheet.title)
                    for i, row in enumerate(sheet.iter_rows(values_only=True)):
                        if i >= 2000:
                            break
                        parts.append("\t".join("" if c is None else str(c) for c in row[:50]))
                text = "\n".join(parts)
            finally:
                book.close()
    elif kind == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted or len(reader.pages) > 50:
            raise ValueError("Encrypted or oversized PDF")
        text = "\n".join(
            f"[Page {i + 1}] " + (p.extract_text() or "") for i, p in enumerate(reader.pages)
        )
    else:
        raise ValueError("Unsupported file type")
    if len(text) > 200000 or not text.strip():
        raise ValueError("Empty or oversized extracted text")
    if SECRET.search(text):
        raise ValueError("Possible secrets rejected")
    return kind, text


def extract_isolated(filename, data):
    environment = dict(os.environ, OPENBLAS_NUM_THREADS="1", OMP_NUM_THREADS="1")
    try:
        response = subprocess.run(
            [sys.executable, "-m", "modax.extract_worker"],
            input=json.dumps(
                {"filename": filename, "data": base64.b64encode(data).decode()}
            ).encode(),
            capture_output=True,
            timeout=15,
            env=environment,
            check=False,
        )
        result = json.loads(response.stdout)
        if response.returncode != 0 or "error" in result:
            raise ValueError("Document parsing failed or exceeded resource limits")
        return result["kind"], result["text"]
    except (subprocess.TimeoutExpired, ValueError, KeyError):
        raise ValueError("Document parsing failed or exceeded resource limits") from None


class UserProviders:
    def __init__(self, store, settings, registry):
        self.store, self.settings, self.registry = store, settings, registry

    def cipher(self):
        if not self.settings.secret_key:
            raise HTTPException(503, "Server secret encryption is not configured")
        return Fernet(self.settings.secret_key.encode())

    def for_owner(self, owner):
        providers = dict(self.registry.providers)
        for row in self.store.query("SELECT * FROM provider_profiles WHERE owner=?", (owner,)):
            try:
                key = self.cipher().decrypt(row["secret"].encode()).decode()
            except Exception:
                continue
            model = Model(row["id"], row["family"], row["model"])
            if row["family"] == "anthropic":
                from .providers import AnthropicProvider

                provider = AnthropicProvider(model, row["base_url"], key)
            else:
                provider = HTTPProvider(model, row["base_url"], key)
            providers[row["id"]] = provider
        return providers

    def select(self, owner, model="auto", capability="text"):
        providers = self.for_owner(owner)
        if model == "auto":
            values = [p for p in providers.values() if capability in p.get_capabilities()]
            values.sort(key=lambda p: p.model.demo)
            if not values:
                raise HTTPException(422, "No configured model for this capability")
            return values[0]
        if model not in providers:
            raise HTTPException(422, "Model is not configured for this account")
        providers[model].require_capability(capability)
        return providers[model]


def preferences(store, owner):
    rows = store.query("SELECT settings FROM user_settings WHERE owner=?", (owner,))
    return (
        json.loads(rows[0]["settings"]) if rows else {"use_memory": False, "requests_per_day": 100}
    )


def context_for(store, owner, ids):
    context = []
    if preferences(store, owner)["use_memory"]:
        context.extend(
            "User preference: " + r["content"]
            for r in store.query(
                "SELECT content FROM memory WHERE owner=? ORDER BY created LIMIT 20", (owner,)
            )
        )
    for file_id in ids:
        try:
            document = store.require("documents", owner, file_id)
        except KeyError:
            raise HTTPException(404, "Attachment not found") from None
        context.append(
            f"[Document {document['id']}: {document['filename']}]\n" + document["content"][:12000]
        )
    return "\n\n".join(context)[:30000]


def install_platform(app):
    store, settings, registry = app.state.store, app.state.settings, app.state.registry
    owner, require = app.state.owner_dependency, app.state.require_resource
    parser_gate = asyncio.Semaphore(1)
    accounts = UserProviders(store, settings, registry)
    app.state.user_providers = accounts

    @app.get("/api/v1/providers")
    def profiles(user=Depends(owner)):
        return store.query(
            "SELECT id,family,model,base_url,created FROM provider_profiles WHERE owner=?", (user,)
        )

    @app.post("/api/v1/providers", status_code=201)
    def add_profile(body: ProfileInput, user=Depends(owner)):
        try:
            parsed = urlparse(body.base_url)
            port = parsed.port
        except ValueError:
            raise HTTPException(422, "Invalid provider URL or port") from None
        if (
            parsed.scheme != "https"
            or parsed.hostname not in ALLOWED_HOSTS
            or parsed.username
            or parsed.password
            or parsed.query
            or parsed.fragment
            or port not in {None, 443}
        ):
            raise HTTPException(422, "Provider URL must use an approved HTTPS endpoint")
        if body.family == "anthropic" and parsed.hostname != "api.anthropic.com":
            raise HTTPException(422, "Anthropic requires its official API endpoint")
        fixed_endpoints = {
            "gemini": "https://generativelanguage.googleapis.com/v1beta/openai",
            "openrouter": "https://openrouter.ai/api/v1",
        }
        if (
            body.family in fixed_endpoints
            and body.base_url.rstrip("/") != fixed_endpoints[body.family]
        ):
            raise HTTPException(422, "Provider requires its official API endpoint")
        if len(profiles(user)) >= 5:
            raise HTTPException(429, "Provider profile quota reached")
        pid = uid()
        encrypted = accounts.cipher().encrypt(body.api_key.encode()).decode()
        store.execute(
            "INSERT INTO provider_profiles VALUES(?,?,?,?,?,?,?)",
            (pid, user, body.family, body.model, body.base_url.rstrip("/"), encrypted, time.time()),
        )
        store.audit(user, "provider.created", {"id": pid, "family": body.family})
        return next(p for p in profiles(user) if p["id"] == pid)

    @app.delete("/api/v1/providers/{profile_id}", status_code=204)
    def delete_profile(profile_id: str, user=Depends(owner)):
        require("provider_profiles", user, profile_id)
        store.execute("DELETE FROM provider_profiles WHERE id=? AND owner=?", (profile_id, user))

    @app.get("/api/v1/settings")
    def get_settings(user=Depends(owner)):
        return preferences(store, user)

    @app.put("/api/v1/settings")
    def set_settings(body: Preferences, user=Depends(owner)):
        store.execute(
            "INSERT INTO user_settings VALUES(?,?) ON CONFLICT(owner) DO UPDATE SET settings=excluded.settings",
            (user, body.model_dump_json()),
        )
        return body.model_dump()

    @app.get("/api/v1/files")
    def files(user=Depends(owner)):
        return store.query(
            "SELECT id,filename,kind,created FROM documents WHERE owner=? ORDER BY created DESC",
            (user,),
        )

    @app.post("/api/v1/files", status_code=201)
    async def upload(file: UploadFile, user=Depends(owner)):
        data = await file.read(200001)
        if len(data) > 200000:
            raise HTTPException(422, "File size limit is 200 KB")
        try:
            async with parser_gate:
                kind, text = await asyncio.to_thread(
                    extract_isolated, file.filename or "unknown", data
                )
        except Exception as exc:
            raise HTTPException(
                422, str(exc) if isinstance(exc, ValueError) else "File could not be parsed"
            ) from None
        if len(files(user)) >= 30:
            raise HTTPException(429, "Document quota reached")
        fid = uid()
        filename = Path((file.filename or "document").replace("\\", "/")).name[:120]
        blocks = [text[i : i + 1500] for i in range(0, len(text), 1300)]
        with store.connect() as db:
            db.execute(
                "INSERT INTO documents VALUES(?,?,?,?,?,?)",
                (fid, user, filename, kind, text, time.time()),
            )
            for i, block in enumerate(blocks):
                db.execute("INSERT INTO chunks VALUES(?,?,?,?)", (uid(), fid, i, block))
        return {
            "id": fid,
            "filename": filename,
            "chunks": len(blocks),
            "characters": len(text),
            "kind": kind,
        }

    @app.get("/api/v1/files/search")
    def search(q: str, user=Depends(owner)):
        if not 1 <= len(q.strip()) <= 100:
            raise HTTPException(422, "Search must be 1..100 characters")
        return store.query(
            "SELECT d.id AS file_id,d.filename,c.ordinal,c.content FROM chunks c JOIN documents d ON c.document_id=d.id WHERE d.owner=? AND c.content LIKE ? LIMIT 20",
            (user, "%" + q + "%"),
        )

    @app.get("/api/v1/files/{file_id}")
    def document(file_id: str, user=Depends(owner)):
        return require("documents", user, file_id)

    @app.delete("/api/v1/files/{file_id}", status_code=204)
    def delete_file(file_id: str, user=Depends(owner)):
        require("documents", user, file_id)
        store.execute("DELETE FROM documents WHERE id=? AND owner=?", (file_id, user))

    @app.get("/api/v1/arena/runs")
    def runs(user=Depends(owner)):
        return [
            json.loads(r["results"])
            for r in store.query(
                "SELECT results FROM evaluations WHERE owner=? ORDER BY created DESC LIMIT 50",
                (user,),
            )
        ]

    @app.post("/api/v1/arena/runs", status_code=201)
    async def evaluate(body: ArenaInput, user=Depends(owner)):
        if "{task}" not in body.template:
            raise HTTPException(422, "Template must contain {task}")
        selected = [accounts.select(user, model) for model in dict.fromkeys(body.models)]
        if any(not p.model.demo for p in selected) and not body.confirm_external_calls:
            raise HTTPException(403, "Confirm external inference calls before a benchmark")
        result = {
            "id": uid(),
            "dataset": "core-v1",
            "created": time.time(),
            "results": [],
            "grading": "exact-match; JSON normalized; small smoke benchmark, not a frontier intelligence evaluation",
        }
        for provider in selected:
            cases = []
            for item in DATASET:
                started = time.perf_counter()
                try:
                    import asyncio

                    async with asyncio.timeout(90):
                        response = await provider.generate_text(
                            [
                                {
                                    "role": "user",
                                    "content": body.template.replace("{task}", item["task"]),
                                }
                            ]
                        )
                    if item["id"] == "json":
                        try:
                            passed = json.loads(response) == json.loads(item["expected"])
                        except ValueError:
                            passed = False
                    else:
                        passed = response.strip() == item["expected"]
                    cases.append(
                        {
                            "id": item["id"],
                            "passed": passed,
                            "response": response[:1000],
                            "latency_ms": round((time.perf_counter() - started) * 1000, 2),
                        }
                    )
                except (ProviderError, TimeoutError):
                    cases.append({"id": item["id"], "error": "Provider failed", "passed": False})
            result["results"].append(
                {
                    "model": provider.model.id,
                    "demo": provider.model.demo,
                    "score": sum(c["passed"] for c in cases) / len(cases),
                    "cases": cases,
                }
            )
        store.execute(
            "INSERT INTO evaluations VALUES(?,?,?,?)",
            (result["id"], user, json.dumps(result), time.time()),
        )
        return result
